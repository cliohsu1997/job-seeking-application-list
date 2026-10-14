"""Compile CICE careers JSON into clean Application-list columns."""
from __future__ import annotations

import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from columns import COLUMNS

from paths import APP_XLSX, DATA_ECON, ensure_dirs

DATA = DATA_ECON

MATERIAL_LABELS = [
    ("Curriculum Vitae (CV)", "CV"),
    ("Curriculum Vitae", "CV"),
    ("CV", "CV"),
    ("Cover Letter", "Cover Letter"),
    ("Job Market Paper", "Job Market Paper"),
    ("Research Papers", "Research Papers"),
    ("Research Paper", "Research Papers"),
    ("Teaching Evaluations", "Teaching Evaluations"),
    ("Teaching Statement", "Teaching Statement"),
    ("Research Statement", "Research Statement"),
    ("Letters of Reference", "Letters of Reference"),
    ("Letter of Recommendation", "Letters of Reference"),
    ("Letters of Recommendation", "Letters of Reference"),
    ("Recommendation Letters", "Letters of Reference"),
    ("Transcripts", "Transcripts"),
]


def clean_text(value: str) -> str:
    value = html.unescape(html.unescape(value or ""))
    value = value.replace("&#32;", " ")
    return re.sub(r"\s+", " ", value).strip()


def extract_required_materials(description: str) -> str | None:
    if not description:
        return None
    found = []
    for label, norm in MATERIAL_LABELS:
        if re.search(rf"\b{re.escape(label)}\b", description, flags=re.I):
            if norm not in found:
                found.append(norm)
    # Common CICE phrasing without explicit list
    if not found and re.search(r"\b(applications?|applicants?)\b", description, flags=re.I):
        if re.search(r"\bPh\.?D\b", description, flags=re.I):
            # typical econ package when ad is generic
            return "CV; Job Market Paper; Letters of Reference"
    return "; ".join(found) if found else None


def item_to_row(item: dict) -> dict:
    url = item.get("url") or ""
    university = clean_text(item.get("institution") or "")
    school = clean_text(item.get("school") or "")
    description = clean_text(item.get("description") or "")
    # legacy fields
    if not school:
        school = clean_text(item.get("title") or "")
        if school.lower() in {"access to details", "to apply for jobs"}:
            school = ""
    if school and university and school.lower() == university.lower():
        school = ""

    materials = extract_required_materials(description)
    email = None
    emails = re.findall(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", description)
    if emails:
        email = "; ".join(list(dict.fromkeys(e.rstrip(".") for e in emails))[:3])

    deadline = item.get("deadline") or item.get("Deadline")
    if deadline in ("", "None"):
        deadline = None

    return {
        "Country": "China",
        "Location": None,
        "University": university or None,
        "School": school or None,
        "Deadline": deadline,
        "Job Post": url or None,  # CICE detail page
        "Required Material": materials,
        "Submission Channel": "CICE",
        "Application Link": url or None,
        "Email": email,
        "Submission Date": None,
        "Reference Status": None,
        "Interview Date": None,
        "Job Type": "Faculty",
        "_title": school or university or "CICE listing",
    }


def load_rows() -> list[dict]:
    files = sorted(DATA.glob("cice_listings_*.json"), reverse=True)
    if not files:
        raise FileNotFoundError("No cice_listings_*.json. Run download_cice.py first.")
    items = json.loads(files[0].read_text(encoding="utf-8"))
    print(f"Compiling {len(items)} CICE listings from {files[0].name}")
    return [item_to_row(item) for item in items]


def _write_sheet(ws, rows: list[dict]) -> None:
    link_font = Font(color="0563C1", underline="single")
    for c, name in enumerate(COLUMNS, 1):
        ws.cell(1, c, name).font = Font(bold=True)
    for r_i, row in enumerate(rows, 2):
        for c, name in enumerate(COLUMNS, 1):
            val = row.get(name)
            cell = ws.cell(r_i, c, val)
            if name == "Job Post" and val:
                cell.value = row.get("_title") or "CICE listing"
                cell.hyperlink = str(val)
                cell.font = link_font
            elif name == "Application Link" and val:
                cell.hyperlink = str(val)
                cell.font = link_font
    for c in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18


def update_application_workbook(rows: list[dict]) -> None:
    wb = load_workbook(APP_XLSX)
    if "CICE Listings" in wb.sheetnames:
        del wb["CICE Listings"]
    ws = wb.create_sheet("CICE Listings")
    _write_sheet(ws, rows)
    wb.save(APP_XLSX)


def main() -> None:
    rows = load_rows()
    update_application_workbook(rows)
    print(f"Compiled {len(rows)} CICE listings")
    print(f"Updated sheet CICE Listings in {APP_XLSX}")


if __name__ == "__main__":
    main()
