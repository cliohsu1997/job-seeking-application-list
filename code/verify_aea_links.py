"""
Check each AEA/JOE listing one-by-one:
  - Job Post  -> working JOE listing page (canonical JOE_ID)
  - Application Link -> real apply URL scraped from that page (verified)

Saves progress so it can resume if interrupted.
"""
from __future__ import annotations

import html as html_lib
import json
import re
import time
from datetime import date
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from columns import COLUMNS
from joe_urls import listing_url, read_year_issue

ROOT = Path(__file__).resolve().parents[1]
from paths import APP_XLSX, DATA_ECON, ensure_dirs

DATA = DATA_ECON
PROGRESS = DATA / "aea_link_check_progress.json"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
SLEEP_SEC = 0.35

SKIP_URL_PARTS = (
    "aeaweb.org",
    "google.com",
    "googletagmanager",
    "recaptcha",
    "facebook.com",
    "twitter.com",
    "linkedin.com",
    "cloudflare",
    "jquery",
    "schema.org",
)


def latest_xml() -> Path:
    files = sorted(DATA.glob("joe_listings_*.xml"), reverse=True)
    if not files:
        raise FileNotFoundError("No joe_listings_*.xml in data/")
    return files[0]


def fetch(url: str, timeout: int = 30) -> tuple[int, str, str]:
    """Return (status, final_url, body_text)."""
    req = Request(url, headers={"User-Agent": UA, "Accept": "text/html"})
    try:
        with urlopen(req, timeout=timeout) as resp:
            body = resp.read().decode("utf-8", "replace")
            return resp.status, resp.geturl(), body
    except HTTPError as e:
        try:
            body = e.read().decode("utf-8", "replace")
        except Exception:
            body = ""
        return e.code, url, body
    except URLError as e:
        return 0, url, str(e.reason)


def is_skipped_url(url: str) -> bool:
    low = url.lower()
    return any(p in low for p in SKIP_URL_PARTS)


def clean_url(url: str) -> str:
    url = html_lib.unescape(url.strip())
    url = url.rstrip(").,;\"'<>")
    return url


def extract_apply_urls(page_html: str) -> list[str]:
    """Pull likely application URLs from a JOE listing HTML page."""
    found: list[str] = []

    def add(u: str) -> None:
        u = clean_url(u)
        if not u.startswith("http"):
            return
        if is_skipped_url(u):
            return
        if u not in found:
            found.append(u)

    # Prefer application-instruction blocks
    blocks = re.findall(
        r"(?:Application Instructions|External Application|To apply|Apply)(.{0,2500})",
        page_html,
        flags=re.I | re.S,
    )
    search_regions = blocks + [page_html]

    for region in search_regions:
        for m in re.finditer(r"""href=["'](https?://[^"']+)["']""", region, flags=re.I):
            add(m.group(1))
        for m in re.finditer(r"(https?://[^\s<>\"']+)", region):
            add(m.group(1))

    # Prefer known ATS hosts
    preferred = []
    other = []
    ats = (
        "interfolio.com",
        "econjobmarket.org",
        "icims.com",
        "workday",
        "greenhouse.io",
        "lever.co",
        "apply.",
        "careers.",
        "jobs.",
        "facultypositions",
    )
    for u in found:
        if any(a in u.lower() for a in ats):
            preferred.append(u)
        else:
            other.append(u)
    return preferred + other


def page_matches_listing(page_html: str, institution: str, title: str) -> bool:
    text = re.sub(r"\s+", " ", html_lib.unescape(page_html)).lower()
    inst = (institution or "").strip().lower()
    tit = (title or "").strip().lower()
    if inst and inst[:20] in text:
        return True
    if tit and len(tit) > 12 and tit[:24] in text:
        return True
    # JOE pages always contain this cycle header when live
    if "joe listings" in text and "application" in text:
        return True
    return False


def check_url_alive(url: str) -> bool:
    status, final, body = fetch(url, timeout=25)
    if status and 200 <= status < 400:
        low = body.lower()
        if "404" in low and "not found" in low and len(body) < 2000:
            return False
        return True
    return False


def load_progress() -> dict:
    if PROGRESS.exists():
        return json.loads(PROGRESS.read_text(encoding="utf-8"))
    return {"checked": {}}


def save_progress(prog: dict) -> None:
    DATA.mkdir(exist_ok=True)
    PROGRESS.write_text(json.dumps(prog, ensure_ascii=False, indent=2), encoding="utf-8")


def iter_positions(xml_path: Path):
    import xml.etree.ElementTree as ET

    year, issue = read_year_issue(xml_path)
    for _event, elem in ET.iterparse(xml_path, events=("end",)):
        tag = elem.tag.split("}")[-1].lower()
        if tag != "position":
            continue
        jp_id = elem.attrib.get("jp_id", "")
        def t(name: str) -> str:
            for c in elem:
                if c.tag.split("}")[-1].lower() == name:
                    return re.sub(r"\s+", " ", "".join(c.itertext())).strip()
            return ""

        countries, places = [], []
        for child in elem:
            if child.tag.split("}")[-1].lower() != "locations":
                continue
            for loc in child:
                if loc.tag.split("}")[-1].lower() != "location":
                    continue
                fields = {
                    c.tag.split("}")[-1].lower(): re.sub(
                        r"\s+", " ", "".join(c.itertext())
                    ).strip()
                    for c in loc
                }
                if fields.get("country") and fields["country"] not in countries:
                    countries.append(fields["country"])
                place = ", ".join(
                    p
                    for p in (
                        fields.get("city", "").title() if fields.get("city") else "",
                        fields.get("state", ""),
                    )
                    if p
                )
                if place and place not in places:
                    places.append(place)

        yield {
            "jp_id": jp_id,
            "year": year,
            "issue": issue,
            "institution": t("jp_institution"),
            "title": t("jp_title"),
            "section": t("jp_section"),
            "deadline": t("jp_application_deadline"),
            "department": t("jp_department") or t("jp_division"),
            "full_text": t("jp_full_text"),
            "country": "; ".join(countries) or None,
            "location": "; ".join(places) or None,
        }
        elem.clear()


def verify_one(pos: dict) -> dict:
    joe_url = listing_url(pos["jp_id"], pos["year"], pos["issue"])
    status, final_url, body = fetch(joe_url)
    ok_listing = bool(status and 200 <= status < 400 and page_matches_listing(
        body, pos["institution"], pos["title"]
    ))
    # Prefer final redirected canonical URL
    listing = final_url.split("#")[0] if ok_listing else joe_url
    if listing.endswith("&"):
        listing = listing[:-1]

    apply_url = None
    apply_ok = False
    candidates = extract_apply_urls(body) if ok_listing else []
    for cand in candidates:
        time.sleep(SLEEP_SEC)
        if check_url_alive(cand):
            apply_url = cand
            apply_ok = True
            break
    if not apply_url:
        # fall back to JOE listing if that is where materials are submitted
        apply_url = listing if ok_listing else None
        apply_ok = ok_listing

    from classify_aea_job_type import classify_joe_fields

    job_type, _reason = classify_joe_fields(
        section=pos.get("section") or "",
        title=pos.get("title") or "",
        full_text=pos.get("full_text") or "",
    )

    return {
        "jp_id": pos["jp_id"],
        "joe_url": listing,
        "joe_ok": ok_listing,
        "joe_status": status,
        "apply_url": apply_url,
        "apply_ok": apply_ok,
        "apply_candidates": candidates[:5],
        "institution": pos["institution"],
        "title": pos["title"],
        "department": pos.get("department"),
        "deadline": pos.get("deadline"),
        "section": pos.get("section"),
        "country": pos.get("country"),
        "location": pos.get("location"),
        "job_type": job_type,
        "full_text": pos.get("full_text") or "",
    }


def rows_from_results(results: list[dict]) -> list[dict]:
    from compile_joe import extract_email, extract_required_materials, parse_deadline

    rows = []
    for r in results:
        joe = r.get("joe_url")
        apply = r.get("apply_url") or joe
        title = r.get("title") or ""
        rows.append(
            {
                "Country": r.get("country"),
                "Location": r.get("location"),
                "University": r.get("institution"),
                "School": r.get("department"),
                "Deadline": parse_deadline(r.get("deadline") or ""),
                "Job Post": joe,
                "Required Material": extract_required_materials(r.get("full_text") or "")
                or None,
                "Submission Channel": "AEA",
                "Application Link": apply,
                "Email": extract_email(r.get("full_text") or "") or None,
                "Submission Date": None,
                "Reference Status": None,
                "Interview Date": None,
                "Job Type": r.get("job_type") or "Faculty",
                "_title": title,
                "_joe_ok": r.get("joe_ok"),
                "_apply_ok": r.get("apply_ok"),
            }
        )
    return rows


def write_verified_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "JOE Listings"
    link_font = Font(color="0563C1", underline="single")
    header_font = Font(bold=True)
    for c, name in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, name)
        cell.font = header_font
    for r_i, row in enumerate(rows, 2):
        for c, name in enumerate(COLUMNS, 1):
            val = row.get(name)
            cell = ws.cell(r_i, c, val)
            if name == "Job Post" and val:
                title = row.get("_title") or "JOE listing"
                cell.value = title
                cell.hyperlink = str(val)
                cell.font = link_font
            elif name == "Application Link" and val:
                cell.value = str(val)
                cell.hyperlink = str(val)
                cell.font = link_font
    for c in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    wb.save(path)


def update_application_sheet(rows: list[dict]) -> None:
    wb = load_workbook(APP_XLSX)
    if "JOE Listings" in wb.sheetnames:
        del wb["JOE Listings"]
    ws = wb.create_sheet("JOE Listings")
    link_font = Font(color="0563C1", underline="single")
    for c, name in enumerate(COLUMNS, 1):
        ws.cell(1, c, name).font = Font(bold=True)
    for r_i, row in enumerate(rows, 2):
        for c, name in enumerate(COLUMNS, 1):
            val = row.get(name)
            cell = ws.cell(r_i, c, val)
            if name == "Job Post" and val:
                cell.value = row.get("_title") or "JOE listing"
                cell.hyperlink = str(val)
                cell.font = link_font
            elif name == "Application Link" and val:
                cell.hyperlink = str(val)
                cell.font = link_font
    wb.save(APP_XLSX)


def main(limit: int | None = None) -> None:
    xml_path = latest_xml()
    year, issue = read_year_issue(xml_path)
    print(f"XML: {xml_path.name}  cycle={year}-{issue}")
    prog = load_progress()
    checked: dict = prog.setdefault("checked", {})

    positions = list(iter_positions(xml_path))
    if limit:
        positions = positions[:limit]
    total = len(positions)
    print(f"Checking {total} listings one-by-one (resume: {len(checked)} saved)")

    results = []
    for i, pos in enumerate(positions, 1):
        jp = pos["jp_id"]
        if jp in checked and checked[jp].get("joe_ok"):
            results.append(checked[jp])
            print(f"[{i}/{total}] skip cached {jp}")
            continue
        print(f"[{i}/{total}] {pos['institution'][:40]} | {pos['title'][:40]}")
        result = verify_one(pos)
        checked[jp] = result
        results.append(result)
        print(
            f"     JOE={'OK' if result['joe_ok'] else 'FAIL'} {result['joe_url']}"
        )
        print(
            f"     APPLY={'OK' if result['apply_ok'] else 'FAIL'} {result['apply_url']}"
        )
        if i % 5 == 0:
            save_progress(prog)
        time.sleep(SLEEP_SEC)

    save_progress(prog)

    rows = rows_from_results(results)
    today = date.today().isoformat()
    out = DATA / f"joe_verified_{today}.xlsx"
    write_verified_xlsx(rows, out)
    update_application_sheet(rows)

    n_joe = sum(1 for r in results if r.get("joe_ok"))
    n_apply = sum(1 for r in results if r.get("apply_ok"))
    n_ext = sum(
        1
        for r in results
        if r.get("apply_url") and "aeaweb.org/joe" not in (r.get("apply_url") or "")
    )
    print(f"Saved {out}")
    print(f"JOE pages OK: {n_joe}/{len(results)}")
    print(f"Apply links OK: {n_apply}/{len(results)} (external apply: {n_ext})")
    print(f"Updated JOE Listings in {APP_XLSX}")


if __name__ == "__main__":
    import sys

    lim = int(sys.argv[1]) if len(sys.argv) > 1 else None
    main(limit=lim)
