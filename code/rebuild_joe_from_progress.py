"""Rebuild JOE Listings from XML compile + any verified apply URLs in progress."""
from __future__ import annotations

import json
from pathlib import Path

from compile_joe import load_joe_rows, main as compile_joe_unused  # noqa: F401
from verify_aea_links import (
    PROGRESS,
    rows_from_results,
    update_application_sheet,
    write_verified_xlsx,
)
from joe_urls import listing_url, read_year_issue
from columns import COLUMNS
from openpyxl.styles import Font
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from datetime import date

ROOT = Path(__file__).resolve().parents[1]
from paths import APP_XLSX, DATA_ECON, ensure_dirs

DATA = DATA_ECON


def main() -> None:
    from compile_joe import (
        extract_email,
        extract_required_materials,
        load_joe_rows,
    )

    xml = sorted(DATA.glob("joe_listings_*.xml"), reverse=True)[0]
    year, issue = read_year_issue(xml)
    compiled = load_joe_rows(xml)

    prog = {}
    if PROGRESS.exists():
        prog = json.loads(PROGRESS.read_text(encoding="utf-8")).get("checked", {})

    rows = []
    n_verified = 0
    for row in compiled:
        # recover jp_id from Job Post URL
        job = str(row.get("Job Post") or "")
        jp_id = ""
        if "_" in job:
            jp_id = job.rsplit("_", 1)[-1].split("&")[0]
        verified = prog.get(jp_id)
        if verified and verified.get("joe_ok"):
            n_verified += 1
            joe = verified.get("joe_url") or job
            apply = verified.get("apply_url") or joe
            row = dict(row)
            row["Job Post"] = joe
            row["Application Link"] = apply
            row["_title"] = verified.get("title") or row.get("_title")
            if verified.get("country"):
                row["Country"] = verified.get("country")
            if verified.get("location"):
                row["Location"] = verified.get("location")
        rows.append(row)

    today = date.today().isoformat()
    out = DATA / f"joe_verified_{today}.xlsx"
    # write with hyperlinks
    wb = Workbook()
    ws = wb.active
    ws.title = "JOE Listings"
    link_font = Font(color="0563C1", underline="single")
    for c, name in enumerate(COLUMNS, 1):
        ws.cell(1, c, name).font = Font(bold=True)
    for r_i, row in enumerate(rows, 2):
        for c, name in enumerate(COLUMNS, 1):
            val = row.get(name)
            cell = ws.cell(r_i, c, val)
            if name == "Job Post" and val:
                cell.value = row.get("_title") or "JOE listing"
                cell.hyperlink = str(val)
                cell.font = link_font
            elif name == "Application Link" and val:
                cell.hyperlink = str(val)
                cell.font = link_font
    for c in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    wb.save(out)

    app = load_workbook(APP_XLSX)
    if "JOE Listings" in app.sheetnames:
        del app["JOE Listings"]
    # copy sheet
    ws2 = app.create_sheet("JOE Listings")
    for row in ws.iter_rows():
        for cell in row:
            n = ws2.cell(cell.row, cell.column, cell.value)
            if cell.hyperlink:
                n.hyperlink = cell.hyperlink.target
                n.font = link_font
    app.save(APP_XLSX)

    print(f"JOE rows: {len(rows)} (apply verified overlay: {n_verified})")
    print(f"Saved {out} and updated Application workbook")


if __name__ == "__main__":
    main()
