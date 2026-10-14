"""
Combine AEA (JOE) + EJM + CICE into clean/economics/ with Application columns.

PDFs are written separately to clean/ (not inside the subfolder).
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from columns import COLUMNS
from paths import APP_XLSX, CLEAN_ECON, ensure_dirs

SOURCE_SHEETS = [
    ("JOE Listings", "AEA"),
    ("EJM Listings", "EJM"),
    ("CICE Listings", "CICE"),
]


def extract_url(text) -> str | None:
    if text is None:
        return None
    s = str(text).strip()
    if not s:
        return None
    if s.startswith("http://") or s.startswith("https://"):
        urls = re.findall(r"https?://[^\s|]+", s)
        return urls[-1].rstrip(").,;\"'") if urls else s
    urls = re.findall(r"https?://[^\s|]+", s)
    return urls[-1].rstrip(").,;\"'") if urls else None


def normalize_row(values: dict, default_channel: str) -> dict:
    row = {c: values.get(c) for c in COLUMNS}
    channel = row.get("Submission Channel") or default_channel
    row["Submission Channel"] = channel

    link = extract_url(row.get("Application Link")) or extract_url(row.get("Job Post"))
    job_post = row.get("Job Post")
    title = None
    if job_post and " | " in str(job_post):
        left, right = str(job_post).rsplit(" | ", 1)
        if extract_url(right):
            title = left.strip() or None
            if not link:
                link = extract_url(right)

    if link:
        row["Application Link"] = link
        if title:
            row["Job Post"] = f"{title} | {link}"
        elif job_post and not extract_url(job_post):
            row["Job Post"] = f"{job_post} | {link}"
        elif not job_post:
            row["Job Post"] = link
        else:
            if link not in str(job_post):
                row["Job Post"] = f"{job_post} | {link}"
    return row


def cell_value_or_link(cell) -> str | None:
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    val = cell.value
    return None if val is None else str(val)


def read_sheet_rows(wb, sheet_name: str, default_channel: str) -> list[dict]:
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        raw = {}
        for c in range(1, len(headers) + 1):
            name = headers[c - 1]
            cell = ws.cell(r, c)
            if name in ("Job Post", "Application Link"):
                raw[name] = cell_value_or_link(cell)
            else:
                raw[name] = cell.value
        if not any(raw.get(h) for h in ("University", "Job Post", "Application Link")):
            continue
        rows.append(normalize_row(raw, default_channel))
    return rows


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "All Listings"
    link_font = Font(color="0563C1", underline="single")
    for c, name in enumerate(COLUMNS, 1):
        ws.cell(1, c, name).font = Font(bold=True)
    for r_i, row in enumerate(rows, 2):
        for c, name in enumerate(COLUMNS, 1):
            val = row.get(name)
            cell = ws.cell(r_i, c, val)
            if name == "Application Link" and val:
                url = extract_url(val)
                if url:
                    cell.hyperlink = url
                    cell.font = link_font
            elif name == "Job Post" and val:
                url = extract_url(val) or extract_url(row.get("Application Link"))
                if url:
                    cell.hyperlink = url
                    cell.font = link_font
    widths = {
        "Country": 12,
        "Location": 22,
        "University": 28,
        "School": 22,
        "Deadline": 12,
        "Job Post": 40,
        "Required Material": 28,
        "Submission Channel": 12,
        "Application Link": 40,
        "Email": 24,
        "Job Type": 12,
    }
    for c, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 14)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def deadline_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if type(value) is date:
        return value
    s = str(value).strip()
    for fmt, n in (("%Y-%m-%d %H:%M:%S", 19), ("%Y-%m-%d", 10), ("%m/%d/%Y", 10)):
        try:
            return datetime.strptime(s[:n], fmt).date()
        except ValueError:
            continue
    return None


def is_outdated(row: dict, today: date) -> bool:
    d = deadline_date(row.get("Deadline"))
    if d is None:
        return False
    return d < today


# Website-channel outputs only — keep provincial / Macau-Taiwan xlsx.
_WEBSITE_XLSX_PREFIXES = (
    "all_listings_",
    "aea_listings_",
    "ejm_listings_",
    "cice_listings_",
)


def clear_old_econ_outputs() -> None:
    if not CLEAN_ECON.exists():
        return
    for p in CLEAN_ECON.glob("*"):
        if not p.is_file():
            continue
        if any(p.name.startswith(prefix) for prefix in _WEBSITE_XLSX_PREFIXES):
            p.unlink()


def main() -> None:
    ensure_dirs()
    clear_old_econ_outputs()

    wb = load_workbook(APP_XLSX, data_only=True)

    all_rows: list[dict] = []
    counts = {}
    for sheet, channel in SOURCE_SHEETS:
        rows = read_sheet_rows(wb, sheet, channel)
        counts[channel] = len(rows)
        all_rows.extend(rows)
    wb.close()

    today_d = date.today()
    with_link = [r for r in all_rows if extract_url(r.get("Application Link"))]
    missing = len(all_rows) - len(with_link)
    active = [r for r in with_link if not is_outdated(r, today_d)]
    outdated = len(with_link) - len(active)

    def job_type_rank(jt: str) -> int:
        if jt in ("Faculty", "Academic"):
            return 0
        if jt == "Postdoc":
            return 1
        return 2

    def sort_key(r):
        d = r.get("Deadline")
        dkey = d if isinstance(d, datetime) else datetime.max
        return (
            r.get("Submission Channel") or "",
            job_type_rank(r.get("Job Type") or ""),
            dkey,
            r.get("University") or "",
        )

    active.sort(key=sort_key)

    today = today_d.isoformat()
    xlsx_path = CLEAN_ECON / f"all_listings_{today}.xlsx"
    write_xlsx(active, xlsx_path)

    for channel in ("AEA", "EJM", "CICE"):
        subset = [r for r in active if r.get("Submission Channel") == channel]
        write_xlsx(subset, CLEAN_ECON / f"{channel.lower()}_listings_{today}.xlsx")

    n_fac = sum(1 for r in active if r.get("Job Type") in ("Faculty", "Academic"))
    n_post = sum(1 for r in active if r.get("Job Type") == "Postdoc")
    n_non = sum(1 for r in active if r.get("Job Type") == "Non-Academic")
    print(f"Sources: {counts}")
    print(f"With Application Link: {len(with_link)}")
    print(f"Removed outdated (deadline < {today}): {outdated}")
    print(f"Clean active listings: {len(active)}")
    print(f"Job Type - Faculty: {n_fac} | Postdoc: {n_post} | Non-Academic: {n_non}")
    if missing:
        print(f"Skipped {missing} rows without a direct link")
    print(f"Wrote {xlsx_path}")
    print(f"Columns: {COLUMNS}")


if __name__ == "__main__":
    main()
